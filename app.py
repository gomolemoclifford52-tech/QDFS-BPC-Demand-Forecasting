import streamlit as st
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import hashlib
import re

from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from forecasting import (
    moving_average_forecast,
    exponential_smoothing_forecast,
    calculate_accuracy
)

import ui


# ============================================================
# 1. PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="QDFS - BPC Demand Forecasting",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown(ui.CSS, unsafe_allow_html=True)


# ============================================================
# 2. DEFAULT BPC PROJECT DATA
# ============================================================

DEFAULT_DATA = pd.DataFrame({
    "Quarter": [
        "2020 Q1", "2020 Q2", "2020 Q3", "2020 Q4",
        "2021 Q1", "2021 Q2", "2021 Q3", "2021 Q4",
        "2022 Q1", "2022 Q2", "2022 Q3", "2022 Q4",
        "2023 Q1", "2023 Q2", "2023 Q3", "2023 Q4",
        "2024 Q1"
    ],
    "Actual_MWh": [
        1011335, 853636, 993552, 983328,
        943147, 968484, 1003738, 1012766,
        1005502, 1099937, 1075525, 1084200,
        1165633, 1124952, 1271961, 1172981,
        1201287
    ]
})


# ============================================================
# 3. HELPER FUNCTIONS
# ============================================================

def make_data_signature(data):
    """Create a stable signature for the loaded dataset."""
    hashed = pd.util.hash_pandas_object(
        data[["Quarter", "Actual_MWh"]],
        index=True
    ).values.tobytes()

    return hashlib.sha256(hashed).hexdigest()


def get_next_quarter_label(last_quarter):
    """Convert labels such as 2024 Q1 into 2024 Q2."""
    match = re.match(
        r"^\s*(\d{4})\s*Q([1-4])\s*$",
        str(last_quarter),
        flags=re.IGNORECASE
    )

    if not match:
        return "Next Quarter"

    year = int(match.group(1))
    quarter = int(match.group(2))

    if quarter == 4:
        year += 1
        quarter = 1
    else:
        quarter += 1

    return f"{year} Q{quarter}"


def build_styled_excel(sheets):
    """Create a formatted Excel workbook from a dictionary of DataFrames."""
    buffer = BytesIO()

    with pd.ExcelWriter(
        buffer,
        engine="openpyxl"
    ) as writer:

        for sheet_name, data in sheets.items():
            data.to_excel(
                writer,
                index=False,
                sheet_name=sheet_name
            )

    buffer.seek(0)

    workbook = load_workbook(buffer)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="243B53"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # Style the header row
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Set reasonable column widths
        for column_cells in worksheet.columns:

            column_letter = get_column_letter(
                column_cells[0].column
            )

            max_length = 0

            for cell in column_cells:
                value = cell.value

                if value is not None:
                    max_length = max(
                        max_length,
                        len(str(value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 3, 42)

        # Apply readable numeric formats
        headers = {
            cell.column: str(cell.value)
            for cell in worksheet[1]
        }

        for row in worksheet.iter_rows(min_row=2):

            for cell in row:

                header = headers.get(
                    cell.column,
                    ""
                ).lower()

                if not isinstance(
                    cell.value,
                    (int, float)
                ):
                    continue

                if "mape" in header or "percentage error" in header:
                    cell.number_format = '0.00"%"'

                elif "alpha" in header:
                    cell.number_format = "0.0"

                elif "mwh" in header or "mad" in header:
                    cell.number_format = "#,##0.00"

                else:
                    cell.number_format = "#,##0.00"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def make_line_chart(
    x,
    series,
    title,
    x_label,
    y_label,
    rotate_x=False,
    marker_x=None,
    marker_label=None
):
    """Create a Matplotlib line chart in the QDFS style."""
    fig, ax = plt.subplots(figsize=(12, 4.4))

    for item in series:
        ax.plot(
            x,
            item["values"],
            marker=item.get("marker", "o"),
            linestyle=item.get("linestyle", "-"),
            linewidth=2.2,
            markersize=5,
            color=item.get("color"),
            label=item["label"]
        )

    if marker_x is not None:
        ax.axvline(
            marker_x,
            color=ui.INK,
            linestyle=(0, (4, 4)),
            linewidth=1.2,
            label=marker_label
        )

    ui.style_axes(fig, ax, title, x_label, y_label)

    if rotate_x:
        plt.setp(
            ax.get_xticklabels(),
            rotation=45,
            ha="right"
        )

    fig.tight_layout()

    return fig


def show_chart(fig):
    st.pyplot(fig, width="stretch")
    plt.close(fig)


def card(name):
    """A styled card. The key gives the container a CSS class."""
    return st.container(key=f"card_{name}")


# ============================================================
# 4. SESSION STATE
# ============================================================

defaults = {
    "forecast_signature": None,
    "forecast_run": False,
    "ma_test_results": None,
    "es_test_results": None,
    "ma_test_data_signature": None,
    "es_test_data_signature": None
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# 5. SIDEBAR: DATA INPUT AND MODEL CONTROLS
# ============================================================

with st.sidebar:

    st.header("Controls")

    st.subheader("1. Data input")

    uploaded_file = st.file_uploader(
        "Upload CSV or Excel",
        type=["csv", "xlsx"],
        help="Required columns: Quarter and Actual_MWh"
    )

    # --------------------------------------------------------
    # Load the dataset
    # --------------------------------------------------------

    if uploaded_file is not None:

        try:
            if uploaded_file.name.lower().endswith(".csv"):
                df = pd.read_csv(uploaded_file)
            else:
                df = pd.read_excel(uploaded_file)

            df = df.dropna(how="all").copy()

            required_columns = [
                "Quarter",
                "Actual_MWh"
            ]

            missing_columns = [
                column
                for column in required_columns
                if column not in df.columns
            ]

            if missing_columns:
                st.error(
                    "Missing required column(s): "
                    + ", ".join(missing_columns)
                )
                st.info(
                    "Your file must contain Quarter and Actual_MWh."
                )
                st.stop()

            df = df[required_columns].copy()

            # Validate before converting to strings
            if df["Quarter"].isna().any():
                st.error("Quarter contains missing values.")
                st.stop()

            # Clean quarter labels
            df["Quarter"] = (
                df["Quarter"]
                .astype(str)
                .str.strip()
            )

            if df["Quarter"].eq("").any():
                st.error("Some Quarter labels are empty.")
                st.stop()

            # Allow numbers formatted with commas
            df["Actual_MWh"] = (
                df["Actual_MWh"]
                .astype(str)
                .str.replace(",", "", regex=False)
                .str.strip()
            )

            df["Actual_MWh"] = pd.to_numeric(
                df["Actual_MWh"],
                errors="coerce"
            )

            if df["Actual_MWh"].isna().any():
                st.error(
                    "Actual_MWh contains missing or non-numeric values."
                )
                st.stop()

            if df["Quarter"].duplicated().any():
                st.error(
                    "Duplicate Quarter labels were found. "
                    "Each quarter should appear only once."
                )
                st.stop()

            if len(df) < 5:
                st.error(
                    "At least 5 observations are required."
                )
                st.stop()

            if (df["Actual_MWh"] <= 0).any():
                st.error(
                    "Actual_MWh values must be greater than zero "
                    "because MAPE requires non-zero actual values."
                )
                st.stop()

            data_source = uploaded_file.name

            st.success(
                f"Loaded {len(df)} observations."
            )

        except Exception as error:
            st.error(
                f"Unable to read the uploaded file: {error}"
            )
            st.stop()

    else:
        df = DEFAULT_DATA.copy()
        data_source = "Default BPC Project Dataset"

        st.info(
            "Using the default BPC project dataset."
        )

    st.caption(
        "Forecasts use rows in the order supplied. "
        "Ensure your data is chronological."
    )

    # --------------------------------------------------------
    # Model parameters
    # --------------------------------------------------------

    st.subheader("2. Model parameters")

    max_window = min(8, len(df) - 1)
    default_window = min(4, max_window)

    ma_window = st.slider(
        "Moving average window",
        min_value=2,
        max_value=max_window,
        value=default_window,
        step=1
    )

    alpha = st.slider(
        "Exponential smoothing α",
        min_value=0.1,
        max_value=0.9,
        value=0.4,
        step=0.1
    )

    # --------------------------------------------------------
    # Signatures and status
    # --------------------------------------------------------

    current_data_signature = make_data_signature(df)

    current_forecast_signature = (
        f"{current_data_signature}|"
        f"{ma_window}|{alpha:.1f}"
    )

    forecast_is_current = (
        st.session_state["forecast_run"]
        and
        st.session_state["forecast_signature"]
        == current_forecast_signature
    )

    st.divider()

    st.subheader("3. Run analysis")

    run_forecast = st.button(
        "Run forecast",
        type="primary",
        width="stretch"
    )

    if run_forecast:
        st.session_state["forecast_run"] = True
        st.session_state["forecast_signature"] = (
            current_forecast_signature
        )
        forecast_is_current = True

    if st.button(
        "Clear forecast",
        width="stretch"
    ):
        st.session_state["forecast_run"] = False
        st.session_state["forecast_signature"] = None
        forecast_is_current = False

    if not st.session_state["forecast_run"]:
        st.info("Forecast not run yet.")

    elif not forecast_is_current:
        st.warning(
            "Data or parameters changed. "
            "Select Run forecast to refresh the results."
        )

    else:
        st.success("Forecast is up to date.")

    st.divider()

    with st.expander("About QDFS"):
        st.write(
            "QDFS is an academic forecasting application "
            "for comparing Moving Average and Exponential "
            "Smoothing on quarterly electricity distribution data."
        )
        st.write(
            "The project uses electricity distribution as a "
            "proxy for demand."
        )


# ============================================================
# 6. TOP BAR AND TABS
# ============================================================

st.html(
    ui.topbar(
        data_source,
        df["Quarter"].iloc[0],
        df["Quarter"].iloc[-1]
    )
)

tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
    "Dashboard",
    "Forecast results",
    "Accuracy",
    "Parameter testing",
    "Data & downloads",
    "Report"
])


# ============================================================
# 7. FORECAST CALCULATIONS
# ============================================================

forecast_available = (
    st.session_state["forecast_run"]
    and
    st.session_state["forecast_signature"]
    == current_forecast_signature
)

forecast_is_stale = (
    st.session_state["forecast_run"]
    and not forecast_available
)

actual_values = df["Actual_MWh"].to_numpy()

next_quarter = get_next_quarter_label(
    df["Quarter"].iloc[-1]
)

forecast_view = None

if forecast_available:

    # --------------------------------------------------------
    # Moving Average
    # --------------------------------------------------------

    ma_forecasts = moving_average_forecast(
        actual_values,
        window=ma_window
    )

    ma_mad, ma_mape = calculate_accuracy(
        actual_values,
        ma_forecasts,
        start_index=ma_window
    )

    # --------------------------------------------------------
    # Exponential Smoothing
    # --------------------------------------------------------

    es_forecasts = exponential_smoothing_forecast(
        actual_values,
        alpha=alpha
    )

    es_mad, es_mape = calculate_accuracy(
        actual_values,
        es_forecasts,
        start_index=1
    )

    # --------------------------------------------------------
    # Historical results
    # --------------------------------------------------------

    results = pd.DataFrame({
        "Quarter": df["Quarter"],
        "Actual_MWh": actual_values,
        "MA_Forecast": ma_forecasts,
        "ES_Forecast": es_forecasts
    })

    results["MA_Error"] = (
        results["Actual_MWh"] - results["MA_Forecast"]
    )

    results["ES_Error"] = (
        results["Actual_MWh"] - results["ES_Forecast"]
    )

    results["MA_Absolute_Error"] = results["MA_Error"].abs()
    results["ES_Absolute_Error"] = results["ES_Error"].abs()

    results["MA_Percentage_Error"] = (
        results["MA_Absolute_Error"]
        / results["Actual_MWh"]
        * 100
    )

    results["ES_Percentage_Error"] = (
        results["ES_Absolute_Error"]
        / results["Actual_MWh"]
        * 100
    )

    # --------------------------------------------------------
    # Next-quarter forecasts
    # --------------------------------------------------------

    next_ma_forecast = np.mean(
        actual_values[-ma_window:]
    )

    next_es_forecast = (
        alpha * actual_values[-1]
        + (1 - alpha) * es_forecasts[-1]
    )

    # --------------------------------------------------------
    # MAD-based ranges
    # --------------------------------------------------------

    ma_lower = next_ma_forecast - ma_mad
    ma_upper = next_ma_forecast + ma_mad

    es_lower = next_es_forecast - es_mad
    es_upper = next_es_forecast + es_mad

    # --------------------------------------------------------
    # Historical MAPE comparison (descriptive only)
    # --------------------------------------------------------

    if ma_mape < es_mape:
        lower_mape_model = f"{ma_window}-Quarter Moving Average"
        lower_mape_short = "Moving average"
        mape_comparison = (
            f"The Moving Average has the lower historical MAPE "
            f"({ma_mape:.2f}% versus {es_mape:.2f}% for Exponential "
            f"Smoothing) under the current settings."
        )
    elif es_mape < ma_mape:
        lower_mape_model = "Exponential Smoothing"
        lower_mape_short = "Exponential smoothing"
        mape_comparison = (
            f"Exponential Smoothing has the lower historical MAPE "
            f"({es_mape:.2f}% versus {ma_mape:.2f}% for the "
            f"Moving Average) under the current settings."
        )
    else:
        lower_mape_model = "Neither model (equal MAPE)"
        lower_mape_short = "Equal"
        mape_comparison = (
            "Both models have the same historical MAPE under the "
            "current settings."
        )

    # --------------------------------------------------------
    # Values for the dashboard components
    # --------------------------------------------------------

    forecast_view = {
        "next_quarter": next_quarter,
        "ma_window": ma_window,
        "alpha": alpha,
        "ma_next": float(next_ma_forecast),
        "es_next": float(next_es_forecast),
        "ma_mad": float(ma_mad),
        "es_mad": float(es_mad),
        "ma_mape": float(ma_mape),
        "es_mape": float(es_mape),
        "ma_lower": float(ma_lower),
        "ma_upper": float(ma_upper),
        "es_lower": float(es_lower),
        "es_upper": float(es_upper),
        "lower_mape_label": lower_mape_short,
        "ma_eval_n": len(df) - ma_window,
        "es_eval_n": len(df) - 1,
        "latest": float(actual_values[-1]),
        "latest_quarter": df["Quarter"].iloc[-1]
    }


# ============================================================
# 8. TAB 1 — DASHBOARD
# ============================================================

with tab1:

    dataset_view = {
        "n": len(df),
        "latest": float(actual_values[-1]),
        "latest_quarter": df["Quarter"].iloc[-1],
        "average": float(np.mean(actual_values)),
        "first_quarter": df["Quarter"].iloc[0],
        "last_quarter": df["Quarter"].iloc[-1]
    }

    st.html(
        ui.hero(
            dataset_view,
            forecast_view,
            next_quarter,
            stale=forecast_is_stale
        )
    )

    st.html(
        ui.dashboard_cards(
            df["Quarter"].tolist(),
            actual_values.tolist(),
            forecast_view
        )
    )

    if forecast_available:

        with card("dashboard_chart"):

            fig = make_line_chart(
                df["Quarter"],
                [
                    {
                        "values": actual_values,
                        "label": "Actual",
                        "color": ui.SERIES_COLORS["actual"]
                    },
                    {
                        "values": ma_forecasts,
                        "label": f"Moving average ({ma_window} quarters)",
                        "marker": "s",
                        "linestyle": "--",
                        "color": ui.SERIES_COLORS["ma"]
                    },
                    {
                        "values": es_forecasts,
                        "label": f"Exponential smoothing (α = {alpha:.1f})",
                        "marker": "^",
                        "linestyle": "--",
                        "color": ui.SERIES_COLORS["es"]
                    }
                ],
                "Actual distribution and historical forecasts",
                "Quarter",
                "Distribution (MWh)",
                rotate_x=True
            )

            show_chart(fig)


# ============================================================
# 9. TAB 2 — FORECAST RESULTS
# ============================================================

with tab2:

    st.html(
        ui.page_header(
            "Forecast results",
            "Next-quarter forecasts from both models, the full "
            "historical forecast table, and forecast errors by quarter."
        )
    )

    if not forecast_available:

        st.info("Run the forecast from the sidebar to see results.")

    else:

        c1, c2 = st.columns(2)

        with c1:
            with card("results_ma"):
                st.subheader("Moving average")
                st.metric(
                    f"{ma_window}-quarter moving average, {next_quarter}",
                    f"{next_ma_forecast:,.0f} MWh"
                )
                st.write(f"**MAD:** {ma_mad:,.2f} MWh")
                st.write(
                    f"**Range:** {ma_lower:,.0f} – "
                    f"{ma_upper:,.0f} MWh"
                )

        with c2:
            with card("results_es"):
                st.subheader("Exponential smoothing")
                st.metric(
                    f"Exponential smoothing (α = {alpha:.1f}), {next_quarter}",
                    f"{next_es_forecast:,.0f} MWh"
                )
                st.write(f"**MAD:** {es_mad:,.2f} MWh")
                st.write(
                    f"**Range:** {es_lower:,.0f} – "
                    f"{es_upper:,.0f} MWh"
                )

        st.caption(
            "The ranges are forecast ± MAD, not confidence intervals."
        )

        with card("results_table"):

            st.subheader("Historical forecast table")

            display_results = results.rename(columns={
                "Actual_MWh": "Actual (MWh)",
                "MA_Forecast": "MA Forecast (MWh)",
                "ES_Forecast": "ES Forecast (MWh)",
                "MA_Error": "MA Error (MWh)",
                "ES_Error": "ES Error (MWh)"
            })

            st.dataframe(
                display_results.style.format({
                    "Actual (MWh)": "{:,.0f}",
                    "MA Forecast (MWh)": "{:,.2f}",
                    "ES Forecast (MWh)": "{:,.2f}",
                    "MA Error (MWh)": "{:,.2f}",
                    "ES Error (MWh)": "{:,.2f}"
                }),
                width="stretch",
                hide_index=True,
                height=500
            )

        with card("results_chart"):

            fig = make_line_chart(
                df["Quarter"],
                [
                    {
                        "values": actual_values,
                        "label": "Actual",
                        "color": ui.SERIES_COLORS["actual"]
                    },
                    {
                        "values": ma_forecasts,
                        "label": f"MA ({ma_window} quarters)",
                        "marker": "s",
                        "linestyle": "--",
                        "color": ui.SERIES_COLORS["ma"]
                    },
                    {
                        "values": es_forecasts,
                        "label": f"ES (α = {alpha:.1f})",
                        "marker": "^",
                        "linestyle": "--",
                        "color": ui.SERIES_COLORS["es"]
                    }
                ],
                "Actual vs historical forecasts",
                "Quarter",
                "Distribution (MWh)",
                rotate_x=True
            )

            show_chart(fig)

        with card("results_errors"):

            fig = make_line_chart(
                df["Quarter"],
                [
                    {
                        "values": results["MA_Absolute_Error"],
                        "label": "Moving average",
                        "color": ui.SERIES_COLORS["ma"]
                    },
                    {
                        "values": results["ES_Absolute_Error"],
                        "label": "Exponential smoothing",
                        "marker": "s",
                        "color": ui.SERIES_COLORS["es"]
                    }
                ],
                "Absolute forecast error",
                "Quarter",
                "Absolute error (MWh)",
                rotate_x=True
            )

            show_chart(fig)


# ============================================================
# 10. TAB 3 — ACCURACY ANALYSIS
# ============================================================

with tab3:

    st.html(
        ui.page_header(
            "Accuracy",
            "Mean absolute deviation and mean absolute percentage "
            "error for each model over its historical forecasts."
        )
    )

    if not forecast_available:

        st.info("Run the forecast from the sidebar to see accuracy.")

    else:

        c1, c2, c3, c4 = st.columns(4)

        c1.metric("MA MAD", f"{ma_mad:,.2f} MWh")
        c2.metric("MA MAPE", f"{ma_mape:.2f}%")
        c3.metric("ES MAD", f"{es_mad:,.2f} MWh")
        c4.metric("ES MAPE", f"{es_mape:.2f}%")

        st.write("")

        with card("accuracy_table"):

            st.subheader("Accuracy summary")

            accuracy_table = pd.DataFrame({
                "Model": [
                    f"{ma_window}-Quarter Moving Average",
                    "Exponential Smoothing"
                ],
                "MAD (MWh)": [ma_mad, es_mad],
                "MAPE (%)": [ma_mape, es_mape]
            })

            st.dataframe(
                accuracy_table.style.format({
                    "MAD (MWh)": "{:,.2f}",
                    "MAPE (%)": "{:.2f}%"
                }),
                width="stretch",
                hide_index=True
            )

            st.caption(
                f"Moving average is scored on {len(df) - ma_window} "
                f"quarters and exponential smoothing on {len(df) - 1}, "
                "because each model starts forecasting at a different point."
            )

        with card("accuracy_chart"):

            fig = make_line_chart(
                df["Quarter"],
                [
                    {
                        "values": results["MA_Percentage_Error"],
                        "label": "Moving average",
                        "color": ui.SERIES_COLORS["ma"]
                    },
                    {
                        "values": results["ES_Percentage_Error"],
                        "label": "Exponential smoothing",
                        "marker": "s",
                        "color": ui.SERIES_COLORS["es"]
                    }
                ],
                "Absolute percentage error by quarter",
                "Quarter",
                "Absolute percentage error (%)",
                rotate_x=True
            )

            show_chart(fig)

        c1, c2 = st.columns(2)

        with c1:
            with card("accuracy_measures"):

                st.subheader("Accuracy measures")

                st.write(
                    "**Mean Absolute Deviation (MAD):** the average "
                    "absolute forecast error, expressed in MWh."
                )

                st.latex(
                    r"MAD = \frac{\sum |A_t-F_t|}{n}"
                )

                st.write(
                    "**Mean Absolute Percentage Error (MAPE):** the "
                    "average absolute forecast error as a percentage "
                    "of the actual value."
                )

                st.latex(
                    r"MAPE = \frac{1}{n}\sum "
                    r"\left|\frac{A_t-F_t}{A_t}\right|\times100"
                )

        # ----------------------------------------------------
        # Project benchmark verification
        # ----------------------------------------------------

        with c2:
            with card("accuracy_benchmark"):

                st.subheader("Project results verification")

                is_default_data = (
                    df.reset_index(drop=True).equals(
                        DEFAULT_DATA.reset_index(drop=True)
                    )
                )

                if (
                    is_default_data
                    and ma_window == 4
                    and abs(alpha - 0.4) < 0.001
                ):

                    benchmark = pd.DataFrame({
                        "Measure": [
                            "Moving Average MAD (MWh)",
                            "Moving Average MAPE (%)",
                            "Exponential Smoothing MAD (MWh)",
                            "Exponential Smoothing MAPE (%)"
                        ],
                        "Project Report": [
                            48056.08,
                            4.28,
                            51187.64,
                            4.86
                        ],
                        "Application": [
                            ma_mad,
                            ma_mape,
                            es_mad,
                            es_mape
                        ]
                    })

                    benchmark["Difference"] = (
                        benchmark["Application"]
                        - benchmark["Project Report"]
                    )

                    st.dataframe(
                        benchmark.style.format({
                            "Project Report": "{:,.2f}",
                            "Application": "{:,.2f}",
                            "Difference": "{:+,.2f}"
                        }),
                        width="stretch",
                        hide_index=True
                    )

                    st.caption(
                        "The project report gives ES MAPE as 4.86%. "
                        "Using full-precision calculations may display "
                        "approximately 4.87%; this small difference can "
                        "arise from rounding during intermediate calculations."
                    )

                else:
                    st.info(
                        "Project benchmark verification is shown when "
                        "the default dataset and project parameters "
                        "(MA window 4, α = 0.4) are selected."
                    )


# ============================================================
# 11. TAB 4 — PARAMETER TESTING
# ============================================================

with tab4:

    st.html(
        ui.page_header(
            "Parameter testing",
            "See how accuracy changes as the window and α vary. "
            "Results stay in place when you switch tabs."
        )
    )

    if not forecast_available:

        st.info(
            "Run the main forecast before running parameter tests."
        )

    else:

        c1, c2 = st.columns(2)

        c1.metric("Selected MA window", f"{ma_window} quarters")
        c2.metric("Selected ES α", f"{alpha:.1f}")

        st.write("")

        # ----------------------------------------------------
        # Moving Average testing
        # ----------------------------------------------------

        with card("param_ma"):

            st.subheader("Moving average window")

            st.write(
                "All windows are evaluated over the same set of "
                "quarters to make the comparison fair."
            )

            if st.button(
                "Test MA windows",
                width="stretch"
            ):

                ma_test_rows = []

                # Use a common evaluation start for every window
                common_start = max_window

                for test_window in range(2, max_window + 1):

                    test_forecasts = moving_average_forecast(
                        actual_values,
                        window=test_window
                    )

                    test_mad, test_mape = calculate_accuracy(
                        actual_values,
                        test_forecasts,
                        start_index=common_start
                    )

                    ma_test_rows.append({
                        "Window": test_window,
                        "MAD (MWh)": test_mad,
                        "MAPE (%)": test_mape
                    })

                st.session_state["ma_test_results"] = pd.DataFrame(
                    ma_test_rows
                )

                st.session_state[
                    "ma_test_data_signature"
                ] = current_data_signature

            ma_test_df = st.session_state["ma_test_results"]

            ma_test_is_current = (
                ma_test_df is not None
                and
                st.session_state["ma_test_data_signature"]
                == current_data_signature
            )

            if ma_test_is_current:

                display_ma = ma_test_df.copy()

                display_ma["Status"] = display_ma["Window"].apply(
                    lambda value:
                    "Selected"
                    if value == ma_window
                    else ""
                )

                st.dataframe(
                    display_ma.style.format({
                        "MAD (MWh)": "{:,.2f}",
                        "MAPE (%)": "{:.2f}%"
                    }),
                    width="stretch",
                    hide_index=True
                )

                c1, c2 = st.columns(2)

                with c1:
                    fig = make_line_chart(
                        ma_test_df["Window"],
                        [{
                            "values": ma_test_df["MAD (MWh)"],
                            "label": "MAD",
                            "color": ui.SERIES_COLORS["ma"]
                        }],
                        "Window vs MAD",
                        "Window (quarters)",
                        "MAD (MWh)",
                        marker_x=ma_window,
                        marker_label=f"Selected window = {ma_window}"
                    )

                    show_chart(fig)

                with c2:
                    fig = make_line_chart(
                        ma_test_df["Window"],
                        [{
                            "values": ma_test_df["MAPE (%)"],
                            "label": "MAPE",
                            "color": ui.SERIES_COLORS["ma"]
                        }],
                        "Window vs MAPE",
                        "Window (quarters)",
                        "MAPE (%)",
                        marker_x=ma_window,
                        marker_label=f"Selected window = {ma_window}"
                    )

                    show_chart(fig)

            else:
                st.info(
                    "Select Test MA windows to generate the comparison."
                )

        # ----------------------------------------------------
        # Exponential Smoothing testing
        # ----------------------------------------------------

        with card("param_es"):

            st.subheader("Exponential smoothing α")

            if st.button(
                "Test α values",
                width="stretch"
            ):

                es_test_rows = []

                for test_alpha in np.round(
                    np.arange(0.1, 1.0, 0.1),
                    1
                ):

                    test_forecasts = exponential_smoothing_forecast(
                        actual_values,
                        alpha=float(test_alpha)
                    )

                    test_mad, test_mape = calculate_accuracy(
                        actual_values,
                        test_forecasts,
                        start_index=1
                    )

                    es_test_rows.append({
                        "Alpha": float(test_alpha),
                        "MAD (MWh)": test_mad,
                        "MAPE (%)": test_mape
                    })

                st.session_state["es_test_results"] = pd.DataFrame(
                    es_test_rows
                )

                st.session_state[
                    "es_test_data_signature"
                ] = current_data_signature

            es_test_df = st.session_state["es_test_results"]

            es_test_is_current = (
                es_test_df is not None
                and
                st.session_state["es_test_data_signature"]
                == current_data_signature
            )

            if es_test_is_current:

                display_es = es_test_df.copy()

                display_es["Status"] = display_es["Alpha"].apply(
                    lambda value:
                    "Selected"
                    if abs(value - alpha) < 0.001
                    else ""
                )

                st.dataframe(
                    display_es.style.format({
                        "Alpha": "{:.1f}",
                        "MAD (MWh)": "{:,.2f}",
                        "MAPE (%)": "{:.2f}%"
                    }),
                    width="stretch",
                    hide_index=True
                )

                c1, c2 = st.columns(2)

                with c1:
                    fig = make_line_chart(
                        es_test_df["Alpha"],
                        [{
                            "values": es_test_df["MAD (MWh)"],
                            "label": "MAD",
                            "color": ui.SERIES_COLORS["es"]
                        }],
                        "α vs MAD",
                        "Alpha (α)",
                        "MAD (MWh)",
                        marker_x=alpha,
                        marker_label=f"Selected α = {alpha:.1f}"
                    )

                    show_chart(fig)

                with c2:
                    fig = make_line_chart(
                        es_test_df["Alpha"],
                        [{
                            "values": es_test_df["MAPE (%)"],
                            "label": "MAPE",
                            "color": ui.SERIES_COLORS["es"]
                        }],
                        "α vs MAPE",
                        "Alpha (α)",
                        "MAPE (%)",
                        marker_x=alpha,
                        marker_label=f"Selected α = {alpha:.1f}"
                    )

                    show_chart(fig)

            else:
                st.info(
                    "Select Test α values to generate the comparison."
                )

    # --------------------------------------------------------
    # Methodology formulas
    # --------------------------------------------------------

    with st.expander("Parameter testing methodology"):

        st.markdown("**Moving average**")

        st.write(
            "The Moving Average forecast uses the previous "
            "n observations:"
        )

        st.latex(
            r"MA_t = \frac{A_{t-1}+A_{t-2}+\cdots+A_{t-n}}{n}"
        )

        st.markdown("**Exponential smoothing**")

        st.write(
            "Exponential Smoothing combines the previous actual "
            "observation with the previous forecast:"
        )

        st.latex(
            r"F_t = \alpha A_{t-1}+(1-\alpha)F_{t-1}"
        )

        st.markdown("**Accuracy measures**")

        st.latex(
            r"MAD = \frac{\sum |A_t-F_t|}{n}"
        )

        st.latex(
            r"MAPE = \frac{1}{n}\sum "
            r"\left|\frac{A_t-F_t}{A_t}\right|\times100"
        )


# ============================================================
# 12. TAB 5 — DATA & DOWNLOADS
# ============================================================

with tab5:

    st.html(
        ui.page_header(
            "Data & downloads",
            "Check the loaded data, then export the historical data, "
            "forecast results, or the complete Excel report."
        )
    )

    validation_df = pd.DataFrame({
        "Validation Check": [
            "Required column: Quarter",
            "Required column: Actual_MWh",
            "At least 5 observations",
            "No missing values",
            "Actual_MWh is numeric",
            "Actual_MWh is greater than zero",
            "Quarter labels are unique"
        ],
        "Status": [
            "Passed",
            "Passed",
            "Passed" if len(df) >= 5 else "Failed",
            "Passed" if not df.isna().any().any() else "Failed",
            "Passed" if pd.api.types.is_numeric_dtype(
                df["Actual_MWh"]
            ) else "Failed",
            "Passed" if (df["Actual_MWh"] > 0).all() else "Failed",
            "Passed" if not df["Quarter"].duplicated().any()
            else "Failed"
        ]
    })

    c1, c2 = st.columns([1, 1.3])

    with c1:
        with card("data_validation"):

            st.subheader("Dataset validation")

            st.dataframe(
                validation_df,
                width="stretch",
                hide_index=True
            )

    with c2:
        with card("data_history"):

            st.subheader("Historical dataset")

            st.caption(f"Data source: {data_source}")

            st.dataframe(
                df.style.format({
                    "Actual_MWh": "{:,.0f}"
                }),
                width="stretch",
                hide_index=True
            )

            st.download_button(
                "Download historical CSV",
                data=df.to_csv(index=False).encode("utf-8"),
                file_name="QDFS_Historical_Data.csv",
                mime="text/csv",
                width="stretch"
            )

    if forecast_available:

        with card("data_summary"):

            st.subheader("Forecast summary")

            forecast_summary = pd.DataFrame({
                "Model": [
                    f"{ma_window}-Quarter Moving Average",
                    "Exponential Smoothing"
                ],
                "Parameter": [
                    f"Window = {ma_window}",
                    f"Alpha = {alpha:.1f}"
                ],
                "Forecast Period": [
                    next_quarter,
                    next_quarter
                ],
                "Next Forecast (MWh)": [
                    next_ma_forecast,
                    next_es_forecast
                ],
                "MAD (MWh)": [
                    ma_mad,
                    es_mad
                ],
                "MAPE (%)": [
                    ma_mape,
                    es_mape
                ],
                "Lower MAD Range (MWh)": [
                    ma_lower,
                    es_lower
                ],
                "Upper MAD Range (MWh)": [
                    ma_upper,
                    es_upper
                ]
            })

            st.dataframe(
                forecast_summary.style.format({
                    "Next Forecast (MWh)": "{:,.2f}",
                    "MAD (MWh)": "{:,.2f}",
                    "MAPE (%)": "{:.2f}%",
                    "Lower MAD Range (MWh)": "{:,.2f}",
                    "Upper MAD Range (MWh)": "{:,.2f}"
                }),
                width="stretch",
                hide_index=True
            )

            accuracy_export = pd.DataFrame({
                "Model": [
                    f"{ma_window}-Quarter Moving Average",
                    "Exponential Smoothing"
                ],
                "Parameter": [
                    f"Window = {ma_window}",
                    f"Alpha = {alpha:.1f}"
                ],
                "MAD (MWh)": [ma_mad, es_mad],
                "MAPE (%)": [ma_mape, es_mape]
            })

            model_configuration = pd.DataFrame({
                "Parameter": [
                    "Data Source",
                    "Number of Observations",
                    "First Quarter",
                    "Last Quarter",
                    "Forecast Period",
                    "Moving Average Window",
                    "Exponential Smoothing Alpha"
                ],
                "Value": [
                    data_source,
                    len(df),
                    df["Quarter"].iloc[0],
                    df["Quarter"].iloc[-1],
                    next_quarter,
                    ma_window,
                    alpha
                ]
            })

            # ------------------------------------------------
            # Include parameter tests if they exist for this data
            # ------------------------------------------------

            ma_tests_valid = (
                st.session_state["ma_test_results"] is not None
                and
                st.session_state["ma_test_data_signature"]
                == current_data_signature
            )

            es_tests_valid = (
                st.session_state["es_test_results"] is not None
                and
                st.session_state["es_test_data_signature"]
                == current_data_signature
            )

            ma_tests_export = (
                st.session_state["ma_test_results"]
                if ma_tests_valid
                else pd.DataFrame({
                    "Note": [
                        "Run MA parameter testing to populate this sheet."
                    ]
                })
            )

            es_tests_export = (
                st.session_state["es_test_results"]
                if es_tests_valid
                else pd.DataFrame({
                    "Note": [
                        "Run ES parameter testing to populate this sheet."
                    ]
                })
            )

            # ------------------------------------------------
            # Build complete Excel report
            # ------------------------------------------------

            sheets = {
                "Historical Data": df,
                "Forecast Results": results,
                "Forecast Summary": forecast_summary,
                "Accuracy Analysis": accuracy_export,
                "Model Configuration": model_configuration,
                "Data Validation": validation_df,
                "MA Parameter Testing": ma_tests_export,
                "ES Parameter Testing": es_tests_export
            }

            excel_report = build_styled_excel(sheets)

            d1, d2 = st.columns(2)

            with d1:
                st.download_button(
                    "Download complete Excel report",
                    data=excel_report,
                    file_name="QDFS_Complete_Forecasting_Report.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    type="primary",
                    width="stretch"
                )

            with d2:
                st.download_button(
                    "Download forecast results CSV",
                    data=results.to_csv(index=False).encode("utf-8"),
                    file_name="QDFS_Forecast_Results.csv",
                    mime="text/csv",
                    width="stretch"
                )

            st.caption(
                "The Excel report includes historical data, forecast "
                "results, forecast summary, accuracy, model settings, "
                "validation, and parameter-testing sheets."
            )

    else:
        st.info(
            "Run the forecast to enable forecast and Excel report downloads."
        )


# ============================================================
# 13. TAB 6 — AUTOMATED FORECASTING REPORT
# ============================================================

with tab6:

    st.html(
        ui.page_header(
            "Report",
            "A written summary of the current run that you can "
            "download as Markdown."
        )
    )

    if not forecast_available:

        st.info("Run the forecast to generate the report.")

    else:

        with card("report_summary"):

            st.subheader("Executive summary")

            st.write(
                f"Using **{len(df)} quarterly observations** from "
                f"**{df['Quarter'].iloc[0]} to {df['Quarter'].iloc[-1]}**, "
                f"QDFS forecasts electricity distribution for **{next_quarter}**. "
                f"The selected Moving Average window is **{ma_window} quarters** "
                f"and the Exponential Smoothing parameter is **α = {alpha:.1f}**."
            )

            c1, c2 = st.columns(2)

            with c1:
                st.metric(
                    f"Moving average, {next_quarter}",
                    f"{next_ma_forecast:,.0f} MWh"
                )
                st.caption(
                    f"MAD: {ma_mad:,.2f} MWh | MAPE: {ma_mape:.2f}%"
                )

            with c2:
                st.metric(
                    f"Exponential smoothing, {next_quarter}",
                    f"{next_es_forecast:,.0f} MWh"
                )
                st.caption(
                    f"MAD: {es_mad:,.2f} MWh | MAPE: {es_mape:.2f}%"
                )

            st.subheader("Historical accuracy comparison")
            st.write(mape_comparison)
            st.caption(
                "This comparison describes the historical evaluation results "
                "for the current dataset and settings. It does not establish "
                "which model will be more accurate in future quarters."
            )

        report_text = f"""# QDFS Automated Forecasting Report

## Dataset and configuration
- Data source: {data_source}
- Observations: {len(df)}
- Historical period: {df['Quarter'].iloc[0]} to {df['Quarter'].iloc[-1]}
- Forecast period: {next_quarter}
- Moving Average window: {ma_window} quarters
- Exponential Smoothing alpha: {alpha:.1f}

## Next-quarter forecasts

| Model | Forecast (MWh) | MAD (MWh) | MAPE (%) | Forecast ± MAD range (MWh) |
|---|---:|---:|---:|---:|
| {ma_window}-Quarter Moving Average | {next_ma_forecast:.2f} | {ma_mad:.2f} | {ma_mape:.2f} | {ma_lower:.2f} – {ma_upper:.2f} |
| Exponential Smoothing | {next_es_forecast:.2f} | {es_mad:.2f} | {es_mape:.2f} | {es_lower:.2f} – {es_upper:.2f} |

## Interpretation
{mape_comparison}

MAD is the mean absolute forecast error in MWh. MAPE is the mean absolute
percentage error. Lower values indicate smaller errors on the evaluated
historical observations, but do not guarantee better future performance.

## Planning considerations
- Treat these results as decision-support estimates, not exact future demand.
- The displayed ranges are forecast ± MAD. They are simple error-based
  ranges, not statistical confidence or prediction intervals.
- This project uses electricity distribution as a proxy for demand.
  Supply constraints or load-shedding may cause distribution to differ
  from underlying electricity demand.
- Check that uploaded observations are in chronological order and that
  the data are appropriate for the planning period.
"""

        with card("report_preview"):

            st.subheader("Report preview")
            st.markdown(report_text)

            st.download_button(
                "Download report (Markdown)",
                data=report_text.encode("utf-8"),
                file_name="QDFS_Automated_Forecasting_Report.md",
                mime="text/markdown",
                type="primary",
                width="stretch"
            )


# ============================================================
# 14. ABOUT QDFS AND PROJECT METHODOLOGY
# ============================================================

st.write("")

with st.expander("About QDFS and the project methodology"):

    st.subheader("Project overview")

    st.write(
        "QDFS is an academic application for comparing "
        "Moving Average and Exponential Smoothing methods "
        "using quarterly electricity distribution data."
    )

    st.write(
        "The project uses electricity distribution as a proxy "
        "for demand. Distribution may not represent all "
        "underlying demand, particularly where supply "
        "constraints or load-shedding affect electricity served."
    )

    st.subheader("Moving average")

    st.latex(
        r"MA_t = \frac{A_{t-1}+A_{t-2}+\cdots+A_{t-n}}{n}"
    )

    st.subheader("Exponential smoothing")

    st.latex(
        r"F_t = \alpha A_{t-1}+(1-\alpha)F_{t-1}"
    )

    st.subheader("Accuracy measures")

    st.latex(
        r"MAD = \frac{\sum |A_t-F_t|}{n}"
    )

    st.latex(
        r"MAPE = \frac{1}{n}\sum "
        r"\left|\frac{A_t-F_t}{A_t}\right|\times100"
    )

    st.subheader("Interpretation of forecast ranges")

    st.info(
        "QDFS displays forecast ± MAD as a simple error-based "
        "range. It is not a statistical confidence interval "
        "or prediction interval."
    )

    st.subheader("Important limitations")

    st.markdown("""
    - Electricity distribution is used as a proxy for demand.
    - Supply constraints can cause observed distribution to differ
      from underlying electricity demand.
    - Forecast accuracy depends on the amount and quality of
      historical data supplied.
    - The application compares Moving Average and Exponential
      Smoothing; it does not establish that either method will
      remain more accurate for future periods.
    """)


# ============================================================
# 15. FOOTER
# ============================================================

st.caption(
    "QDFS, Quantitative Demand Forecasting System. "
    "Botswana Power Corporation academic project. "
    "Methods: Moving Average and Exponential Smoothing. "
    "Accuracy: MAD and MAPE."
)
